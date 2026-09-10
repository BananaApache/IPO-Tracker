"""
Shared pipeline for the Wikipedia / IPO-underpricing analysis.

Tests the statement:
    "Companies that already had a Wikipedia article before going public
     are more underpriced."

Sample window is 2014-2024, chosen because free price data does not cover
firms that have since delisted. The Wikipedia match is automated on the
company's own name plus redirect/rename handling; harder relationships
(parent, subsidiary, spinoff, product) are escalated for human review
rather than guessed.

All network access is to free, unauthenticated endpoints. No API keys.
"""

from __future__ import annotations

import json
import os
import re
import time
import threading
import urllib.error
import urllib.parse
import urllib.request
from html import unescape
from pathlib import Path

# Code lives in research/underpricing/; its data lives alongside the rest of
# the study's outputs, under research/data/underpricing/.
PKG = Path(__file__).resolve().parent
ROOT = PKG.parent / "data" / "underpricing"
RAW = ROOT / "raw"
CACHE = ROOT / "cache"
OUT = ROOT / "out"
for _d in (RAW, CACHE, OUT):
    _d.mkdir(parents=True, exist_ok=True)

# SEC requires a descriptive UA with contact info; Wikipedia and Yahoo are
# happier with one too.
CONTACT = "Daniel Li daniel.li@miami.edu"
UA_SEC = f"{CONTACT} (IPO research project)"
UA_WEB = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"
UA_WIKI = f"IPOWikipediaThesis/1.0 ({CONTACT}) python-urllib"


# --------------------------------------------------------------------------
# rate limiting + fetch
# --------------------------------------------------------------------------
class Throttle:
    """Token-bucket-ish limiter shared across threads."""

    def __init__(self, per_sec: float):
        self.interval = 1.0 / per_sec
        self.lock = threading.Lock()
        self.next_at = 0.0

    def wait(self):
        with self.lock:
            now = time.monotonic()
            if now < self.next_at:
                time.sleep(self.next_at - now)
                now = time.monotonic()
            self.next_at = max(now, self.next_at) + self.interval


SEC_LIMIT = Throttle(8)     # SEC asks for <= 10 req/s
WIKI_LIMIT = Throttle(5)
YF_LIMIT = Throttle(6)


def fetch(url, ua=UA_SEC, limiter=None, byte_range=None, tries=3, timeout=45):
    """GET with retries. Returns bytes, or None on persistent failure."""
    headers = {"User-Agent": ua, "Accept-Encoding": "gzip, deflate"}
    if byte_range:
        headers["Range"] = f"bytes=0-{byte_range}"
    last = None
    for attempt in range(tries):
        if limiter:
            limiter.wait()
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                data = r.read()
            if r.headers.get("Content-Encoding") == "gzip":
                import gzip
                data = gzip.decompress(data)
            return data
        except urllib.error.HTTPError as e:
            last = e
            if e.code in (403, 429, 503):
                time.sleep(1.5 * (attempt + 1))
                continue
            if e.code == 404:
                return None
        except Exception as e:  # noqa: BLE001 - network is best-effort
            last = e
            time.sleep(1.0 * (attempt + 1))
    return None


def fetch_json(url, **kw):
    b = fetch(url, **kw)
    if not b:
        return None
    try:
        return json.loads(b.decode("utf8", "ignore"))
    except Exception:  # noqa: BLE001
        return None


# --------------------------------------------------------------------------
# disk cache
# --------------------------------------------------------------------------
def _cache_path(stage: str, key: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", key)[:120]
    d = CACHE / stage
    d.mkdir(parents=True, exist_ok=True)
    return d / f"{safe}.json"


def cached(stage: str, key: str, fn, force=False):
    """Run fn() once per key, persisting the result as JSON."""
    p = _cache_path(stage, key)
    if p.exists() and not force:
        try:
            return json.loads(p.read_text())
        except Exception:  # noqa: BLE001 - corrupt cache entry, recompute
            pass
    val = fn()
    try:
        p.write_text(json.dumps(val))
    except Exception:  # noqa: BLE001
        pass
    return val


# --------------------------------------------------------------------------
# 1. universe (Ritter)
# --------------------------------------------------------------------------
LEGAL_SUFFIXES = r"""
    inc|inc\.|incorporated|corp|corp\.|corporation|co|co\.|company|
    ltd|ltd\.|limited|llc|l\.l\.c\.|lp|l\.p\.|plc|p\.l\.c\.|
    holdings?|holding|group|groups|
    sa|s\.a\.|nv|n\.v\.|ab|as|a\/s|oyj|spa|s\.p\.a\.|ag|se|bv|b\.v\.|
    trust|reit|partners
"""


def normalize_name(name: str, drop_generic=True) -> str:
    """Company name -> a form that has a chance of matching a Wikipedia title."""
    s = (name or "").strip()
    s = re.sub(r"\s*\(.*?\)\s*", " ", s)
    s = s.replace("&", " and ")
    s = re.sub(r"[,\.]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    if drop_generic:
        pat = re.compile(rf"\s+(?:{LEGAL_SUFFIXES})\s*$", re.I | re.X)
        prev = None
        while prev != s:
            prev = s
            cut = pat.sub("", s).strip()
            # Keep the suffix if removing it leaves too little to match on.
            if len(cut) >= 3 and len(cut.split()) >= 1 and len(cut) >= 0.34 * len(s):
                s = cut
    return s


def _is_spac(name: str, ticker: str) -> bool:
    t = str(ticker or "")
    if (t.endswith("U") and len(t) >= 4) or ".U" in t:
        return True
    n = (name or "").upper()
    return bool(re.search(r"ACQUISITION|BLANK CHECK|ACQ\s+CORP", n))


def load_universe(year_min=2014, year_max=2024):
    """Ritter IPO-age.xlsx -> list of IPO dicts, screened."""
    import openpyxl

    wb = openpyxl.load_workbook(RAW / "IPO-age.xlsx", read_only=True)
    ws = wb[wb.sheetnames[0]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        raw_date, name, ticker, cusip, adr, vc, dual, shares, internet, perm, founding = r[:11]
        if raw_date is None:
            continue
        try:
            ds = str(int(raw_date))
            year = int(ds[:4])
        except (TypeError, ValueError):
            continue
        if not (year_min <= year <= year_max):
            continue
        if adr == 2:                       # drop ADRs
            continue
        ticker = str(ticker) if ticker is not None else ""
        name = str(name or "").strip()
        if not ticker or ticker == "." or not name:
            continue
        if _is_spac(name, ticker):         # drop blank-check / unit offerings
            continue
        out.append({
            "ipo_date": f"{ds[:4]}-{ds[4:6]}-{ds[6:8]}",
            "year": year,
            "name": name,
            "ticker": ticker.replace(".U", "").strip(),
            "cusip": str(cusip) if cusip else None,
            "vc": int(vc) if vc in (0, 1) else None,
            "founding": founding if isinstance(founding, int) and founding > 1500 else None,
        })
    return out


# --------------------------------------------------------------------------
# 2. CIK resolution
# --------------------------------------------------------------------------
_CIK_BY_TICKER = None
_CIK_BY_NAME = None


def _load_cik_tables():
    """company_tickers.json (current filers) + cik-lookup-data.txt (all history)."""
    global _CIK_BY_TICKER, _CIK_BY_NAME
    if _CIK_BY_TICKER is not None:
        return
    _CIK_BY_TICKER = {}
    j = json.loads((RAW / "company_tickers.json").read_text())
    for v in j.values():
        _CIK_BY_TICKER[v["ticker"].upper()] = str(v["cik_str"]).zfill(10)

    _CIK_BY_NAME = {}
    # Format: "COMPANY NAME:0000012345:" — includes former names, so this is
    # what rescues firms that were renamed or have since delisted.
    with open(RAW / "cik-lookup-data.txt", "r", encoding="latin-1") as fh:
        for line in fh:
            line = line.rstrip("\n")
            if not line.endswith(":"):
                continue
            parts = line.rsplit(":", 2)
            if len(parts) < 2:
                continue
            nm, cik = parts[0], parts[1]
            if not cik.isdigit():
                continue
            key = _name_key(nm)
            if key and key not in _CIK_BY_NAME:
                _CIK_BY_NAME[key] = cik.zfill(10)


def _name_key(nm: str) -> str:
    s = normalize_name(nm, drop_generic=True).upper()
    return re.sub(r"[^A-Z0-9]", "", s)


def resolve_cik(ipo: dict):
    """Ticker first (fast, exact), then historical-name lookup."""
    _load_cik_tables()
    t = ipo["ticker"].upper()
    if t in _CIK_BY_TICKER:
        return _CIK_BY_TICKER[t], "ticker"
    k = _name_key(ipo["name"])
    if k and k in _CIK_BY_NAME:
        return _CIK_BY_NAME[k], "name"
    # last resort: first two words of the name
    words = normalize_name(ipo["name"]).split()
    if len(words) >= 2:
        k2 = re.sub(r"[^A-Z0-9]", "", " ".join(words[:2]).upper())
        if k2 in _CIK_BY_NAME:
            return _CIK_BY_NAME[k2], "name2"
    return None, "miss"


# --------------------------------------------------------------------------
# 3. offer price from the 424B prospectus cover
# --------------------------------------------------------------------------
_DATE_RE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")


def _daynum(d: str) -> int:
    m = _DATE_RE.match(d)
    y, mo, dd = int(m.group(1)), int(m.group(2)), int(m.group(3))
    return y * 372 + mo * 31 + dd


PRICE_PATTERNS = [
    r"initial\s+public\s+offering\s+price\s+(?:is|of|per\s+share\s+is)?[^$\d]{0,120}\$\s*(\d{1,3}(?:\.\d{2})?)",
    r"price\s+to\s+(?:the\s+)?public[^$]{0,160}?\$\s*(\d{1,3}(?:\.\d{2})?)",
    r"\$\s*(\d{1,3}\.\d{2})\s*per\s+share",
    r"offering\s+price\s+of\s+\$\s*(\d{1,3}(?:\.\d{2})?)\s*per\s+share",
]
SHARES_PATTERN = re.compile(r"([\d][\d,]{5,})\s+[Ss]hares", re.M)


def find_prospectus(cik: str, ipo_date: str, window=25):
    """Locate the 424B filing closest to the offer date."""
    sub = fetch_json(f"https://data.sec.gov/submissions/CIK{cik}.json",
                     ua=UA_SEC, limiter=SEC_LIMIT)
    if not sub:
        return None
    rec = sub.get("filings", {}).get("recent", {})
    blocks = [rec]
    # "recent" caps at 1000 filings; older ones live in separate JSON chunks.
    for extra in sub.get("filings", {}).get("files", []) or []:
        nm = extra.get("name")
        if not nm:
            continue
        older = fetch_json(f"https://data.sec.gov/submissions/{nm}",
                           ua=UA_SEC, limiter=SEC_LIMIT)
        if older:
            blocks.append(older)

    best = None
    target = _daynum(ipo_date)
    rows = []
    for blk in blocks:
        rows.extend(zip(blk.get("form", []), blk.get("filingDate", []),
                        blk.get("accessionNumber", []),
                        blk.get("primaryDocument", [])))
    for form, fdate, acc, doc in rows:
        if not form.startswith("424B"):
            continue
        try:
            dist = abs(_daynum(fdate) - target)
        except Exception:  # noqa: BLE001
            continue
        if dist <= window and (best is None or dist < best[0]):
            best = (dist, form, fdate, acc, doc)
    if not best:
        return None
    _, form, fdate, acc, doc = best
    url = (f"https://www.sec.gov/Archives/edgar/data/{int(cik)}/"
           f"{acc.replace('-', '')}/{doc}")
    return {"form": form, "filed": fdate, "url": url,
            "sic": sub.get("sic"), "sic_desc": sub.get("sicDescription")}


def parse_offer(url: str):
    """Offer price + shares offered from the prospectus cover page."""
    # The cover page is at the front; 900KB is plenty and saves ~GBs overall.
    raw = fetch(url, ua=UA_SEC, limiter=SEC_LIMIT, byte_range=900_000)
    if not raw:
        return None, None
    txt = unescape(re.sub(r"<[^>]+>", " ", raw.decode("utf8", "ignore")))
    txt = re.sub(r"\s+", " ", txt)
    head = txt[:150_000]
    price = None
    for pat in PRICE_PATTERNS:
        m = re.search(pat, head, re.I)
        if m:
            try:
                v = float(m.group(1))
            except ValueError:
                continue
            if 1.0 <= v <= 500.0:
                price = v
                break
    shares = None
    m = SHARES_PATTERN.search(head[:20_000])
    if m:
        try:
            n = int(m.group(1).replace(",", ""))
            if 100_000 <= n <= 5_000_000_000:
                shares = n
        except ValueError:
            pass
    return price, shares


# --------------------------------------------------------------------------
# 4. first-day close (Yahoo)
# --------------------------------------------------------------------------
def first_day_close(ticker: str, ipo_date: str):
    """
    Close on the first session at/after the offer date.

    Yahoo's `close` is split-adjusted, so any split after the IPO would
    distort it; we un-adjust using the split events the API returns.
    """
    url = ("https://query1.finance.yahoo.com/v8/finance/chart/"
           f"{urllib.parse.quote(ticker)}?period1=0&period2=9999999999"
           "&interval=1d&events=split")
    j = fetch_json(url, ua=UA_WEB, limiter=YF_LIMIT)
    try:
        res = j["chart"]["result"][0]
    except (TypeError, KeyError, IndexError):
        return None
    meta = res.get("meta", {})
    if meta.get("instrumentType") != "EQUITY":
        return None
    ts = res.get("timestamp") or []
    quote = (res.get("indicators", {}).get("quote") or [{}])[0]
    closes = quote.get("close") or []
    if not ts or not closes:
        return None

    target = time.mktime(time.strptime(ipo_date, "%Y-%m-%d"))
    pick_t = pick_c = None
    for t, c in zip(ts, closes):
        if c is None:
            continue
        if t >= target - 2 * 86400:
            pick_t, pick_c = t, c
            break
    if pick_c is None:
        return None
    # Must actually be the IPO week, not a series that starts years later.
    if abs(pick_t - target) > 7 * 86400:
        return None

    # Un-adjust for splits that happened after the IPO.
    factor = 1.0
    for k, v in (res.get("events", {}).get("splits") or {}).items():
        try:
            if int(k) > pick_t:
                factor *= float(v["numerator"]) / float(v["denominator"])
        except Exception:  # noqa: BLE001
            continue
    return {"close": pick_c * factor, "close_adj": pick_c,
            "split_factor": factor, "bar_ts": pick_t}


# --------------------------------------------------------------------------
# 5. Wikipedia matcher
# --------------------------------------------------------------------------
WAPI = "https://en.wikipedia.org/w/api.php?"

COMPANY_HINTS = re.compile(
    r"\b(company|corporation|firm|business|manufacturer|retailer|bank|"
    r"airline|brand|subsidiary|conglomerate|startup|enterprise|holdings?|"
    r"provider|operator|developer|publisher|chain|insurer|producer)\b", re.I)
CATEGORY_HINTS = re.compile(
    r"(companies|corporations|businesses|banks|airlines|manufacturers|"
    r"brands|retailers|initial public offerings)", re.I)


def _wapi(**params):
    params.setdefault("format", "json")
    params.setdefault("formatversion", "2")
    return fetch_json(WAPI + urllib.parse.urlencode(params),
                      ua=UA_WIKI, limiter=WIKI_LIMIT, tries=4)


def wiki_candidates(name: str, max_cand=4):
    """(candidates, api_ok). api_ok=False means the lookup failed, not that
    the company has no article."""
    cands, seen = [], set()
    api_ok = False

    def add(t, src):
        if t and t not in seen:
            seen.add(t)
            cands.append({"title": t, "via": src})

    norm = normalize_name(name)
    variants = [norm]
    if norm.lower().endswith("s") and "'" not in norm:
        variants.append(norm[:-1] + "'s")          # Leslies -> Leslie's
    # Direct title hit, following redirects — this is what catches renames
    # (e.g. "Liberty Oilfield Services" -> "Liberty Energy").
    probes = list(dict.fromkeys(variants + [name]))
    for probe in probes:
        j = _wapi(action="query", titles=probe, redirects=1, prop="info")
        if not j:
            continue
        api_ok = True
        for pg in j.get("query", {}).get("pages", []):
            if not pg.get("missing"):
                add(pg.get("title"), "title")
    # Search fallback
    for v in variants:
        j = _wapi(action="query", list="search", srsearch=v, srlimit=max_cand)
        if j:
            api_ok = True
            for hit in j.get("query", {}).get("search", []):
                add(hit["title"], "search")
    return cands[:max_cand + 3], api_ok


def _toks(s: str):
    s = normalize_name(s).lower().replace("'", "").replace("\u2019", "")
    return {w for w in re.findall(r"[a-z0-9]{3,}", s)}


def _token_overlap(a: str, b: str) -> float:
    """Symmetric: penalises candidates carrying extra words of their own."""
    ta, tb = _toks(a), _toks(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / max(len(ta), len(tb))


def looks_like_company(title: str, extract: str, categories) -> bool:
    cats = " ".join(categories or [])
    if CATEGORY_HINTS.search(cats):
        return True
    if COMPANY_HINTS.search((extract or "")[:400]):
        return True
    return False


def revision_at(title: str, when: str):
    """Wikitext of the revision live on `when`, plus the creation timestamp."""
    j = _wapi(action="query", prop="revisions", titles=title,
              rvstart=when + "T23:59:59Z", rvdir="older", rvlimit=1,
              rvprop="ids|timestamp|content", rvslots="main", redirects=1)
    if j is None:
        return None
    try:
        pg = j["query"]["pages"][0]
    except (TypeError, KeyError, IndexError):
        return {"revid": None, "timestamp": None, "wikitext": ""}
    if pg.get("missing") or "revisions" not in pg:
        return {"revid": None, "timestamp": None, "wikitext": ""}
    rev = pg["revisions"][0]
    return {"revid": rev.get("revid"), "timestamp": rev.get("timestamp"),
            "wikitext": rev.get("slots", {}).get("main", {}).get("content", "")}


def creation_ts(title: str):
    """(timestamp, api_ok). api_ok=False means the lookup failed."""
    j = _wapi(action="query", prop="revisions", titles=title,
              rvdir="newer", rvlimit=1, rvprop="timestamp", redirects=1)
    if j is None:
        return None, False
    try:
        return j["query"]["pages"][0]["revisions"][0]["timestamp"], True
    except (TypeError, KeyError, IndexError):
        return None, True


def body_word_count(wikitext: str):
    """Main-body word count, for the 30-word minimum-content rule."""
    if not wikitext:
        return 0, False
    is_redirect = bool(re.match(r"\s*#\s*REDIRECT", wikitext, re.I))
    t = wikitext
    t = re.sub(r"<ref[^>]*>.*?</ref>|<ref[^>]*/>", " ", t, flags=re.S | re.I)
    t = re.sub(r"<!--.*?-->", " ", t, flags=re.S)
    for _ in range(4):  # nested templates / infoboxes
        t2 = re.sub(r"\{\{[^{}]*\}\}", " ", t)
        if t2 == t:
            break
        t = t2
    t = re.sub(r"\[\[(?:File|Image|Category):[^\]]*\]\]", " ", t, flags=re.I)
    t = re.sub(r"\[\[(?:[^\]|]*\|)?([^\]]*)\]\]", r"\1", t)
    t = re.sub(r"^\s*[\*#:;].*$", " ", t, flags=re.M)
    t = re.sub(r"^\s*={2,}.*$", " ", t, flags=re.M)
    t = re.sub(r"<[^>]+>", " ", t)
    t = re.sub(r"'{2,}", " ", t)
    return len(re.findall(r"[A-Za-z][A-Za-z'\-]*", t)), is_redirect


MIN_BODY_WORDS = 30
MIN_OVERLAP = 0.67


def wikipedia_flag(name: str, ipo_date: str, debug=False):
    """
    Returns a dict with the Wikipedia indicator and the evidence behind it.

    Indicator is 1 when an article about the company existed before the IPO
    date with at least 30 words of main body.
    """
    result = {"wiki": 0, "title": None, "created": None, "words": 0,
              "redirect": False, "overlap": None, "via": None,
              "reason": "no_candidate", "ambiguous": False,
              "status": "reject", "extract": "", "would_be": None}

    cands, api_ok = wiki_candidates(name)
    if not api_ok:
        result["wiki"] = None
        result["reason"] = "api_error"
        return result
    if not cands:
        return result

    # Fetch metadata for all candidates in one call.
    titles = "|".join(c["title"] for c in cands)
    meta = _wapi(action="query", titles=titles, redirects=1,
                 prop="extracts|categories", exintro=1, explaintext=1,
                 exlimit="max", cllimit="max")
    info = {}
    if meta:
        for pg in meta.get("query", {}).get("pages", []):
            info[pg.get("title")] = {
                "extract": pg.get("extract", ""),
                "cats": [c.get("title", "") for c in pg.get("categories", [])],
            }
        # follow redirect mapping so candidate titles line up with results
        for rd in meta.get("query", {}).get("redirects", []) or []:
            if rd.get("to") in info:
                info[rd.get("from")] = info[rd["to"]]

    scored = []
    for c in cands:
        t = c["title"]
        i = info.get(t, {})
        ov = _token_overlap(name, t)
        company = looks_like_company(t, i.get("extract", ""), i.get("cats"))
        # Redirect-resolved exact title hits are trustworthy even when the
        # destination name shares no tokens (Nu Holdings -> Nubank).
        score = ov + (0.5 if company else 0) + (0.6 if c["via"] == "title" else 0)
        scored.append({**c, "overlap": ov, "company": company, "score": score,
                       "extract": i.get("extract", "")[:300]})
    scored.sort(key=lambda x: -x["score"])
    best = scored[0]

    strong = best["company"] and best["overlap"] >= MIN_OVERLAP
    exact = best["via"] == "title" and best["overlap"] >= 0.9
    accept = strong or exact
    grey = (not accept) and best["company"] and \
           (best["via"] == "title" or best["overlap"] >= 0.15)

    result["overlap"] = round(best["overlap"], 3)
    result["via"] = best["via"]
    result["title"] = best["title"]
    result["extract"] = best.get("extract", "")[:240]
    result["ambiguous"] = bool(grey)
    result["status"] = "accept" if accept else ("ambiguous" if grey else "reject")
    if debug:
        result["candidates"] = scored

    if not accept:
        result["reason"] = "ambiguous_needs_review" if grey else "candidate_rejected"
        # Still gather the dating evidence so a reviewer can decide quickly.
        if grey:
            created, ok = creation_ts(best["title"])
            result["created"] = created
            if ok and created and created[:10] <= ipo_date:
                rev = revision_at(best["title"], ipo_date)
                if rev:
                    w, rd = body_word_count(rev.get("wikitext", ""))
                    result["words"] = w
                    result["redirect"] = rd
                    result["would_be"] = 0 if (rd or w < MIN_BODY_WORDS) else 1
                else:
                    result["would_be"] = 0
            else:
                result["would_be"] = 0
        return result

    created, ok = creation_ts(best["title"])
    result["created"] = created
    if not ok:
        result["wiki"] = None
        result["reason"] = "api_error"
        return result
    if not created or created[:10] > ipo_date:
        result["reason"] = "created_after_ipo"
        return result

    rev = revision_at(best["title"], ipo_date)
    if rev is None:
        result["wiki"] = None
        result["reason"] = "api_error"
        return result
    if not rev.get("wikitext") and rev.get("revid") is None:
        result["reason"] = "no_revision_at_ipo"
        return result
    words, is_redirect = body_word_count(rev["wikitext"])
    result["words"] = words
    result["redirect"] = is_redirect
    result["revid"] = rev["revid"]
    if is_redirect:
        result["reason"] = "redirect_stub"
        return result
    if words < MIN_BODY_WORDS:
        result["reason"] = "under_30_words"
        return result
    result["wiki"] = 1
    result["reason"] = "ok"
    result["status"] = "accept"
    if debug:
        result["candidates"] = scored
    return result


# --------------------------------------------------------------------------
# benchmark set: firms whose correct coding is known independently, used as
# regression tests on the matcher
# --------------------------------------------------------------------------
BENCHMARK_CASES = [
    # (company name to search, IPO date, expected indicator, note)
    ("LinkedIn", "2011-05-19", 1, "long-standing article well before listing"),
    ("Allot Communications", "2006-11-15", 0, "redirect to 'Allotment' at IPO"),
    ("Zoetis", "2013-02-01", 1, "article created on its IPO date"),
    ("Veridian", "2002-06-13", 0, "disambiguation page, not an article"),
    ("The Hertz Corporation", "2006-11-16", 1, "subsidiary of the listing entity"),
    ("New York Mercantile Exchange", "2006-11-17", 1, "subsidiary of the listing entity"),
]


# --------------------------------------------------------------------------
# adjudicating the ambiguous class
# --------------------------------------------------------------------------
# A redirect can land on a genuine rename or subsidiary (Nu Holdings ->
# Nubank) or on an unrelated firm that
# held the name in a different era (Array Technologies -> ATI Technologies:
# ATI used that name for four months in 1985 and was bought by AMD in 2006,
# fourteen years before the solar-tracker firm IPO'd).
#
# Two free signals separate them:
#   1. industry alignment  - EDGAR gives us the IPO firm's SIC description
#   2. already-defunct     - an entity acquired/dissolved before the IPO date
#                            cannot be the firm that is going public

STOP = set("""the of and or a an inc corp co ltd llc services service etc
nec other misc miscellaneous products product general goods related devices
equipment except industries industry manufacturing operators""".split())

DEFUNCT_RE = re.compile(
    r"\b(was|were)\s+(a|an)\b|\bdefunct\b|\bacquired by\b|\bmerged (in)?to\b|"
    r"\bdissolved\b|\bceased operations\b|\bformer(ly)?\b", re.I)
YEAR_RE = re.compile(r"\b(19|20)\d{2}\b")


def _stem(t: str) -> str:
    """Crude prefix stem so finance/financial and bank/banking collide."""
    return t[:6]


def _industry_tokens(text: str):
    toks = re.findall(r"[a-z]{4,}", (text or "").lower())
    return {_stem(t) for t in toks if t not in STOP}


def industry_alignment(sic_desc: str, extract: str, categories=None) -> float:
    """
    Share of SIC-description stems present in the article.

    Substring containment, not token equality, so "Oil & Gas Field Services"
    still matches an article that says "oilfield services".
    """
    sic_t = _industry_tokens(sic_desc)
    if not sic_t:
        return 0.0
    blob = ((extract or "") + " " + " ".join(categories or [])).lower()
    blob_stems = _industry_tokens(blob)
    hit = sum(1 for t in sic_t if t in blob_stems or t in blob)
    return hit / len(sic_t)


def defunct_before(extract: str, ipo_date: str) -> bool:
    """True when the article describes an entity that had already ended."""
    if not extract or not DEFUNCT_RE.search(extract):
        return False
    ipo_year = int(ipo_date[:4])
    years = [int(y) for y in YEAR_RE.findall(extract)]
    # An acquisition/closure year comfortably before the IPO is the giveaway.
    return any(y < ipo_year - 1 for y in years) if years else False


def refine_ambiguous(row: dict):
    """
    Verdict for one ambiguous row: 1, 0, or None if it still needs a human.

    Returns (verdict, reason, signals).
    """
    extract = row.get("wiki_extract") or ""
    sic_desc = row.get("sic_desc") or ""
    ipo = row["ipo_date"]
    align = industry_alignment(sic_desc, extract)
    dead = defunct_before(extract, ipo)
    would = row.get("wiki_would_be")
    sig = {"industry_alignment": round(align, 3), "defunct_before_ipo": dead,
           "overlap": row.get("wiki_overlap")}

    if dead and align < 0.34:
        return 0, "article subject predates/ended before this IPO", sig
    if align >= 0.34 and not dead:
        return would, "industry matches the filer's SIC", sig
    if (row.get("wiki_overlap") or 0) >= 0.5 and not dead:
        return would, "strong name overlap, entity still active", sig
    return None, "needs human review", sig


# Ground truth confirmed by hand during this build. Nu Holdings/Nubank is a
# genuine parent/subsidiary pair; Array Technologies/ATI is a name collision
# across eras and not the same company at all.
ADJUDICATED_CASES = [
    {"name": "Nu Holdings", "ipo_date": "2021-12-09", "truth": 1,
     "why": "Nu Holdings Ltd. (NYSE: NU) is the parent company of Nubank "
            "-- genuine parent/subsidiary match"},
    {"name": "Array Technologies", "ipo_date": "2020-10-15", "truth": 0,
     "why": "ATI held the name 'Array Technologies Inc.' for four months in "
            "1985 and was acquired by AMD in 2006 -- unrelated to the 2020 "
            "solar-tracker listing"},
]
